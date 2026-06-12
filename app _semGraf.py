import os
from datetime import datetime

import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =====================================================
# CONFIGURAÇÕES
# =====================================================

OUTPUT_DIR = "output"
os.makedirs(OUTPUT_DIR, exist_ok=True)

st.set_page_config(page_title="Relatório SENAI / Claro", layout="wide")


# =====================================================
# ORDEM FIXA DO RELATÓRIO
# =====================================================

ORDEM_RELATORIO = [
    {
        "key": "servicos_consumo",
        "titulo": "Serviços de Campo – Consumo e PME",
        "unidade": "CONSUMO E PME",
        "grupo_atividade": "Serviços de Campo",
        "tem_atividade": True,
        "tem_despesa": True,
    },
    {
        "key": "manutencao_consumo",
        "titulo": "Manutenção de Clientes – Consumo e PME",
        "unidade": "CONSUMO E PME",
        "grupo_atividade": "Manutenção de Clientes",
        "tem_atividade": True,
        "tem_despesa": True,
    },
    {
        "key": "instalacao_consumo",
        "titulo": "Instalação de Clientes e Swap - Consumo e PME",
        "unidade": "CONSUMO E PME",
        "grupo_atividade": "Instalação de Clientes e Swap",
        "tem_atividade": True,
        "tem_despesa": True,
    },
    {
        "key": "servicos_tecnicos_consumo",
        "titulo": "Serviços Técnicos - Consumo e PME",
        "unidade": "CONSUMO E PME",
        "grupo_atividade": "Serviços Técnicos",
        "tem_atividade": False,
        "tem_despesa": True,
    },
    {
        "key": "servicos_empresarial",
        "titulo": "Serviços de Campo – Empresarial",
        "unidade": "EMPRESARIAL",
        "grupo_atividade": "Serviços de Campo",
        "tem_atividade": True,
        "tem_despesa": True,
    },
    {
        "key": "manutencao_empresarial",
        "titulo": "Manutenção de Clientes - Empresarial",
        "unidade": "EMPRESARIAL",
        "grupo_atividade": "Manutenção de Clientes",
        "tem_atividade": True,
        "tem_despesa": True,
    },
    {
        "key": "instalacao_empresarial",
        "titulo": "Instalação e Swap - Empresarial",
        "unidade": "EMPRESARIAL",
        "grupo_atividade": "Instalação de Clientes e Swap",
        "tem_atividade": True,
        "tem_despesa": True,
    },
    {
        "key": "servicos_tecnicos_empresarial",
        "titulo": "Serviços Técnicos - Empresarial",
        "unidade": "EMPRESARIAL",
        "grupo_atividade": "Serviços Técnicos",
        "tem_atividade": False,
        "tem_despesa": True,
    },
    {
        "key": "projetos_empresarial",
        "titulo": "Projetos Especiais - Empresarial",
        "unidade": "EMPRESARIAL",
        "grupo_atividade": "Projetos especiais",
        "tem_atividade": True,
        "tem_despesa": False,
    },
]


# =====================================================
# ORDEM FIXA DOS SUBGRUPOS
# =====================================================

ORDEM_SUBGRUPOS = [
    "PRÉ-VT",
    "Manutenção de Clientes",
    "COP",
    "Manutenção de MDU",
    "Instalação de Clientes",
    "Swap",
    "COP Adesão",
    "Serviços Técnicos",
    "Projetos Especiais",
]


# =====================================================
# FORMATAÇÕES
# =====================================================

def formatar_moeda(valor):
    try:
        valor = float(valor)
    except Exception:
        valor = 0

    sinal = "+" if valor > 0 else "-" if valor < 0 else ""
    valor_abs = abs(valor)

    if valor_abs >= 1000:
        texto = f"{sinal}R${valor_abs / 1000:,.1f}K"
    else:
        texto = f"{sinal}R${valor_abs:,.2f}"

    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def formatar_pct(valor):
    try:
        valor = float(valor)
    except Exception:
        valor = 0

    if abs(valor) <= 1:
        valor = valor * 100

    sinal = "+" if valor > 0 else "-" if valor < 0 else ""
    return f"{sinal}{abs(valor):.1f}%".replace(".", ",")


def tipo_movimento(valor):
    try:
        valor = float(valor)
    except Exception:
        valor = 0

    if valor > 0:
        return "aumento"
    if valor < 0:
        return "redução"
    return "variação"


def seta_movimento(valor):
    try:
        valor = float(valor)
    except Exception:
        valor = 0

    if valor > 0:
        return "↑"
    if valor < 0:
        return "↓"
    return "•"


def normalizar(texto):
    return str(texto).strip().upper()


# =====================================================
# LEITURA DO EXCEL
# =====================================================

def carregar_excel(arquivo):
    df = pd.read_excel(arquivo)
    df.columns = [str(c).strip() for c in df.columns]

    colunas_necessarias = [
        "UNIDADE",
        "Grupo atividades",
        "Variação Grupo atividades",
        "% variação",
        "Descrição Atividade",
        "Descrição Cod",
        "Variação",
        "%",
        "Justificativa",
    ]

    faltantes = [c for c in colunas_necessarias if c not in df.columns]

    if faltantes:
        st.error(f"Colunas faltando no Excel: {faltantes}")
        st.stop()

    base = pd.DataFrame()

    base["unidade"] = df["UNIDADE"]
    base["grupo_atividade"] = df["Grupo atividades"]
    base["variacao_grupo"] = df["Variação Grupo atividades"]
    base["percentual_grupo"] = df["% variação"]

    base["subgrupo"] = df["Descrição Atividade"]
    base["despesa"] = df["Descrição Cod"]
    base["variacao"] = df["Variação"]
    base["percentual"] = df["%"]
    base["justificativa"] = df["Justificativa"]

    if "Mes" in df.columns:
        base["mes"] = df["Mes"]
    elif "Mês" in df.columns:
        base["mes"] = df["Mês"]
    else:
        base["mes"] = "Abr"

    for col in ["variacao_grupo", "percentual_grupo", "variacao", "percentual"]:
        base[col] = pd.to_numeric(base[col], errors="coerce").fillna(0)

    for col in ["unidade", "grupo_atividade", "subgrupo", "despesa", "justificativa", "mes"]:
        base[col] = base[col].fillna("").astype(str).str.strip()

    return base


def obter_mes_relatorio(df):
    if df is None or df.empty:
        return "Abril"

    meses = df["mes"].dropna().astype(str).str.strip()
    meses = meses[meses != ""]

    if meses.empty:
        return "Abril"

    mes = meses.iloc[0]

    meses_extenso = {
        "Jan": "Janeiro",
        "Fev": "Fevereiro",
        "Mar": "Março",
        "Abr": "Abril",
        "Mai": "Maio",
        "Jun": "Junho",
        "Jul": "Julho",
        "Ago": "Agosto",
        "Set": "Setembro",
        "Out": "Outubro",
        "Nov": "Novembro",
        "Dez": "Dezembro",
    }

    return meses_extenso.get(mes, mes)


# =====================================================
# FILTROS E ORDEM
# =====================================================

def filtrar_bloco(df, unidade, grupo_atividade):
    base = df.copy()

    base = base[
        base["unidade"].str.upper().str.contains(unidade.upper(), na=False)
    ]

    base = base[
        base["grupo_atividade"].str.upper().str.contains(grupo_atividade.upper(), na=False)
    ]

    return base


def ordenar_subgrupos(df_subgrupos):
    ordem = {normalizar(nome): i for i, nome in enumerate(ORDEM_SUBGRUPOS)}

    df_subgrupos["ordem_modelo"] = df_subgrupos["subgrupo"].apply(
        lambda x: ordem.get(normalizar(x), 999)
    )

    df_subgrupos["abs_variacao"] = df_subgrupos["variacao"].abs()

    df_subgrupos = df_subgrupos.sort_values(
        by=["ordem_modelo", "abs_variacao"],
        ascending=[True, False]
    )

    return df_subgrupos


# =====================================================
# TEXTO AUTOMÁTICO DO RELATÓRIO
# =====================================================

def montar_texto_bloco(df, unidade, grupo_atividade, titulo):
    base = filtrar_bloco(df, unidade, grupo_atividade)

    if base.empty:
        return f"""{titulo}: Impacto de R$0,00 (0,0%)

Influenciadores:

*Sem grandes variações no período."""

    grupos_unicos = (
        base[["grupo_atividade", "variacao_grupo", "percentual_grupo"]]
        .drop_duplicates()
    )

    variacao_total = grupos_unicos["variacao_grupo"].sum()
    percentual_total = grupos_unicos["percentual_grupo"].sum()

    linhas = [
        f"{titulo}: Impacto de {formatar_moeda(variacao_total)} ({formatar_pct(percentual_total)})",
        "",
        "Influenciadores:",
        "",
    ]

    subgrupos = (
        base.groupby("subgrupo", dropna=False)
        .agg(
            variacao=("variacao", "sum"),
            percentual=("percentual", "sum"),
        )
        .reset_index()
    )

    subgrupos = ordenar_subgrupos(subgrupos)

    for _, sub in subgrupos.iterrows():
        nome_subgrupo = str(sub["subgrupo"]).strip() or "Outros"
        var_subgrupo = sub["variacao"]
        pct_subgrupo = sub["percentual"]

        linhas.append(
            f"{nome_subgrupo}: Impacto de {formatar_moeda(var_subgrupo)} ({formatar_pct(pct_subgrupo)})"
        )

        despesas = base[base["subgrupo"] == nome_subgrupo].copy()

        despesas = despesas[despesas["variacao"] != 0]
        despesas["abs_variacao"] = despesas["variacao"].abs()
        despesas = despesas.sort_values("abs_variacao", ascending=False)

        for _, row in despesas.iterrows():
            despesa = str(row["despesa"]).strip()
            justificativa = str(row["justificativa"]).strip()
            variacao = row["variacao"]
            percentual = row["percentual"]

            linhas.append(
                f"{seta_movimento(variacao)}\t{despesa}: "
                f"{tipo_movimento(variacao)} de {formatar_moeda(variacao)} "
                f"({formatar_pct(percentual)})"
            )

            if justificativa:
                linhas.append(f"•\t{justificativa}")

        linhas.append("")

    return "\n".join(linhas)


def montar_textos_paginas(df):
    textos = {}

    for bloco in ORDEM_RELATORIO:
        textos[bloco["key"]] = montar_texto_bloco(
            df=df,
            unidade=bloco["unidade"],
            grupo_atividade=bloco["grupo_atividade"],
            titulo=bloco["titulo"],
        )

    return textos


# =====================================================
# WORD - FORMATAÇÃO
# =====================================================

def configurar_pagina(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.4)


def add_cabecalho(doc, mes_relatorio):
    doc.add_paragraph("")

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_after = Pt(0)

    r = p.add_run("SENAI – Unidade São Paulo")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(90, 90, 90)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_after = Pt(18)

    r = p.add_run("Sumário – ")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(90, 90, 90)

    r = p.add_run(str(mes_relatorio))
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 64, 128)
    r.underline = True

    r = p.add_run(" de 2026")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(90, 90, 90)


def nova_pagina(doc, mes_relatorio):
    doc.add_page_break()
    add_cabecalho(doc, mes_relatorio)


def add_texto(doc, texto, tamanho=9, bold=False, cor=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

    r = p.add_run(texto)
    r.font.name = "Arial"
    r.font.size = Pt(tamanho)
    r.bold = bold

    if cor:
        r.font.color.rgb = cor

    return p


def add_linha_influenciador(doc, linha):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

    cor_vermelha = RGBColor(255, 0, 0)
    cor_azul = RGBColor(0, 112, 192)
    tamanho = Pt(9)

    if linha.startswith("↑"):
        r = p.add_run("↑ ")
        r.font.name = "Arial"
        r.font.size = tamanho
        r.font.color.rgb = cor_vermelha

        r = p.add_run(linha[1:].strip())
        r.font.name = "Arial"
        r.font.size = tamanho
        r.bold = True
        r.font.color.rgb = cor_vermelha

    elif linha.startswith("↓"):
        r = p.add_run("↓ ")
        r.font.name = "Arial"
        r.font.size = tamanho
        r.font.color.rgb = cor_azul

        r = p.add_run(linha[1:].strip())
        r.font.name = "Arial"
        r.font.size = tamanho
        r.bold = True
        r.font.color.rgb = cor_azul

    else:
        r = p.add_run(linha)
        r.font.name = "Arial"
        r.font.size = tamanho


def add_secao(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)

    r = p.add_run(texto)
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.bold = True
    r.underline = True
    r.font.color.rgb = RGBColor(0, 64, 96)

    pPr = p._element.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F4E79")

    border.append(bottom)
    pPr.append(border)


def add_linha_sumario(doc, titulo, pagina, nivel=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45 * nivel)
    p.paragraph_format.space_after = Pt(0)

    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(
        Cm(17.2),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )

    r = p.add_run(titulo)
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.bold = nivel > 0

    r = p.add_run("\t")
    r.font.name = "Arial"
    r.font.size = Pt(8)

    r = p.add_run(str(pagina))
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.bold = nivel > 0


# =====================================================
# PÁGINAS DO WORD
# =====================================================

def pagina_sumario(doc, mes_relatorio):
    add_cabecalho(doc, mes_relatorio)

    add_texto(doc, "SUMÁRIO EXECUTIVO", 8)
    doc.add_paragraph("")

    linhas = [
        ("Serviços - Consumo", "2", 0),
        ("I.    POR ATIVIDADE", "2", 1),
        ("II.   POR DESPESA", "3", 1),
        ("Manutenção - Consumo", "4", 0),
        ("I.    POR ATIVIDADE", "5", 1),
        ("II.   POR DESPESA", "6", 1),
        ("Instalação dep - Consumo", "7", 0),
        ("I.    POR ATIVIDADE", "7", 1),
        ("II.   POR DESPESA", "8", 1),
        ("Serviços dep - Consumo", "9", 0),
        ("I.    POR DESPESA", "9", 1),
        ("Serviços Camp – Empresarial", "11", 0),
        ("I.    POR ATIVIDADE", "11", 1),
        ("II.   POR DESPESA", "11", 1),
        ("Manutenção- Empresarial", "12", 0),
        ("I.    POR ATIVIDADE", "12", 1),
        ("II.   POR DESPESA", "12", 1),
        ("Instalação - Empresarial", "14", 0),
        ("I.    POR ATIVIDADE", "14", 1),
        ("II.   POR DESPESA", "14", 1),
        ("Serviços Técnicos - Empresarial", "16", 0),
        ("I.    POR DESPESA", "16", 1),
        ("Projetos - Empresarial", "17", 0),
        ("I.    POR ATIVIDADE", "17", 1),
    ]

    for titulo, pagina, nivel in linhas:
        add_linha_sumario(doc, titulo, pagina, nivel)


def pagina_bloco(doc, mes_relatorio, bloco, texto):
    nova_pagina(doc, mes_relatorio)

    add_texto(doc, bloco["titulo"], 9, True)

    if texto:
        for linha in texto.split("\n"):
            add_linha_influenciador(doc, linha)

    if bloco["tem_atividade"]:
        doc.add_paragraph("")
        add_secao(doc, "I.   Por Atividade")

    if bloco["tem_despesa"]:
        doc.add_paragraph("")
        add_secao(doc, "II.   Por Despesa")


def gerar_docx(textos_paginas, mes_relatorio):
    doc = Document()
    configurar_pagina(doc)

    pagina_sumario(doc, mes_relatorio)

    for bloco in ORDEM_RELATORIO:
        texto = textos_paginas.get(bloco["key"], "")
        pagina_bloco(doc, mes_relatorio, bloco, texto)

    caminho = os.path.join(
        OUTPUT_DIR,
        f"relatorio_senai_sem_graficos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
    )

    doc.save(caminho)
    return caminho


# =====================================================
# STREAMLIT
# =====================================================

st.title("Relatório SENAI / Claro — Sem Gráficos")

st.write(
    "Este app gera o relatório em DOCX sem gráficos, respeitando a ordem fixa do modelo "
    "e a ordem dos subgrupos."
)

arquivo_excel = st.file_uploader(
    "Importar Excel do relatório",
    type=["xlsx", "xls"],
)

if arquivo_excel:
    df = carregar_excel(arquivo_excel)
    mes_relatorio = obter_mes_relatorio(df)
    textos_paginas = montar_textos_paginas(df)

    st.success("Excel carregado com sucesso.")
    st.info(f"Mês identificado: {mes_relatorio}")

    with st.expander("Prévia dos dados tratados"):
        st.dataframe(df, use_container_width=True)

    st.divider()

    st.markdown("## Prévia do relatório")

    st.markdown("### Página 1 — Sumário Executivo")

    st.text("""
SUMÁRIO EXECUTIVO

Serviços - Consumo................................................2
   I.    POR ATIVIDADE.............................................2
   II.   POR DESPESA...............................................3
Manutenção - Consumo...............................................4
   I.    POR ATIVIDADE.............................................5
   II.   POR DESPESA...............................................6
Instalação dep - Consumo...........................................7
   I.    POR ATIVIDADE.............................................7
   II.   POR DESPESA...............................................8
Serviços dep - Consumo.............................................9
   I.    POR DESPESA...............................................9
Serviços Camp – Empresarial.......................................11
   I.    POR ATIVIDADE............................................11
   II.   POR DESPESA..............................................11
Manutenção- Empresarial...........................................12
   I.    POR ATIVIDADE............................................12
   II.   POR DESPESA..............................................12
Instalação - Empresarial..........................................14
   I.    POR ATIVIDADE............................................14
   II.   POR DESPESA..............................................14
Serviços Técnicos - Empresarial...................................16
   I.    POR DESPESA..............................................16
Projetos - Empresarial............................................17
   I.    POR ATIVIDADE............................................17
""")

    for bloco in ORDEM_RELATORIO:
        st.divider()
        st.markdown(f"### {bloco['titulo']}")
        st.text(textos_paginas.get(bloco["key"], ""))

        if bloco["tem_atividade"]:
            st.markdown("#### I. Por Atividade")

        if bloco["tem_despesa"]:
            st.markdown("#### II. Por Despesa")

    st.divider()

    if st.button("Gerar DOCX"):
        caminho = gerar_docx(textos_paginas, mes_relatorio)

        with open(caminho, "rb") as f:
            st.download_button(
                "Baixar DOCX",
                data=f,
                file_name=os.path.basename(caminho),
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

else:
    st.info("Importe o Excel para começar.")